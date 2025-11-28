#!/usr/bin/env python3
"""
Batch experiment runner for CDAC and Capacitance/Shape experiments
Supports multiple models via --prefix and --model-name parameters

Usage:
    # CDAC experiments (default)
    python run_batch_experiments.py --prefix claude --model-name claude --template-type array
    python run_batch_experiments.py --prefix claude --model-name claude --template-type full
    
    # Capacitance/Shape experiments (19 capacitances × 5 shapes = 95 experiments)
    python run_batch_experiments.py --experiment-type capacitance_shape --prefix claude --model-name claude
    
    # Preview generated prompt
    python run_batch_experiments.py --prefix claude --model-name claude --template-type full --preview-prompt first
    
    # Run with specific range
    python run_batch_experiments.py --prefix claude --start-index 1 --stop-index 5

Template Configuration:
    To add a new prompt template, add an entry to PROMPT_TEMPLATES dictionary below:
    
    Example:
        "template_name": {
            "description": "Description of what this template does",
            "key_suffix": "suffix_for_prompt_keys",
            "template": "Template content with variables..."
        }
    
    Available template variables (CDAC templates):
        - {excel_file}: Path to Excel file
        - {sheet_name}: Sheet name
        - {safe_sheet_name}: Sheet name (safe for Virtuoso)
        - {cell_prefix}: Prefix for cell names (e.g., 'claude_')
        - {key_suffix}: Key suffix from this template config
    
    Available template variables (capacitance_shape template):
        - {cap_value}: Capacitance value in fF
        - {cap_str}: Capacitance string (e.g., "0_1" for 0.1, "1" for 1)
        - {shape_key}: Shape key (e.g., "H", "H_shieldless")
        - {shape_desc}: Shape description
        - {task_desc}: Task description
"""

# ============================================================================
# Prompt Templates Registry
# ============================================================================
# Add new templates here. Each template must have:
#   - description: Human-readable description
#   - key_suffix: Suffix used in prompt key generation (e.g., "full_flow", "only_array")
#   - template: Template string with placeholders: {excel_file}, {sheet_name}, 
#               {safe_sheet_name}, {cell_prefix}, {key_suffix}
# ============================================================================

PROMPT_TEMPLATES = {
    "full": {
        "description": "Full flow (unit + dummy + array)",
        "key_suffix": "full_flow",
        "template": """Task: Design a complete CDAC capacitor system in three phases:
1. Phase 1: Design unit H-shape capacitor with target capacitance of 6 fF
2. Phase 2: Generate dummy capacitor based on the unit capacitor
3. Phase 3: Generate capacitor array using the unit and dummy capacitors

Design constraints:
- Target capacitance: 6 fF
- Height constraint: less than 2 µm
- Low parasitic capacitance to ground required
- Top plates connected together, placed on outside; bottom plates in middle for adjacency

Configuration:
- Technology: 28nm process node
- Library: LLM_Layout_Design
- Cell names: C_MAIN_6fF_{cell_prefix}{key_suffix}_{safe_sheet_name} (unit), C_DUMMY_6fF_{cell_prefix}{key_suffix}_{safe_sheet_name} (dummy), C_CDAC_6fF_{cell_prefix}{key_suffix}_{safe_sheet_name} (array)
- CDAC array layout file: {excel_file}, Sheet name: {sheet_name}.

Phase 1 - Unit Capacitor:
1. Generate unit H-shape capacitor SKILL code
2. Execute the generated SKILL code using run_il_file or run_il_with_screenshot tool
3. Run DRC verification using run_drc tool on the unit cell
4. Run PEX extraction using run_pex tool on the unit cell
5. Verify capacitance meets target (6 fF) and iterate if necessary

Phase 2 - Dummy Capacitor:
1. Generate dummy capacitor SKILL code based on the unit capacitor
2. Execute the generated SKILL code
3. Run DRC verification on the dummy cell
4. Run PEX extraction on the dummy cell

Phase 3 - CDAC Array:
1. Generate CDAC array SKILL code using Python for calculations, output results as hard-coded values in SKILL
2. Execute the generated SKILL code using run_il_file or run_il_with_screenshot tool
3. Run DRC verification using run_drc tool on the generated array cell
4. Run PEX extraction using run_pex tool on the generated array cell
5. Verify the results and iterate if necessary (fix errors, regenerate if DRC fails, etc.)

"""
    },
    "array": {
        "description": "Array only (skip unit and dummy)",
        "key_suffix": "only_array",
        "template": """Task: Generate CDAC array (skip unit cell and dummy generation, but MUST execute and verify the array).

Initial setup (required):
- Load knowledge base at the beginning
- Perform Information Retrieval (IR) to gather relevant design information
- Use retrieved knowledge to inform CDAC array generation

Unit capacitor specifications (pre-designed):
- Shape: H-shape capacitor
- Dimensions: unit_height = 1.920um, unit_width = 1.430um
- Frame: frame_horizontal_width = 0.110um, frame_vertical_width = 0.110um
- Metal layers: M7 M6 M5 M4 M3
- BOT terminal offset: bot_offset_x = 0um, bot_offset_y = 0um

Configuration:
- Technology: 28nm process node
- Library: LLM_Layout_Design
- Cell names: C_MAIN_6fF (unit), C_DUMMY_6fF (dummy), C_CDAC_6fF_{cell_prefix}{key_suffix}_{safe_sheet_name} (array)
- CDAC array layout file: {excel_file}, Sheet name: {sheet_name}.

Instructions:
- Load knowledge base and perform IR at the start
- Skip unit cell generation and dummy capacitor generation (assume they already exist)
- Proceed directly to CDAC array generation

MANDATORY execution steps (DO NOT SKIP):
1. Generate CDAC array SKILL code using Python for calculations, output results as hard-coded values in SKILL
2. Execute the generated SKILL code using run_il_file or run_il_with_screenshot tool
3. Run DRC verification using run_drc tool on the generated array cell
4. Run PEX extraction using run_pex tool on the generated array cell
5. Verify the results and iterate if necessary (fix errors, regenerate if DRC fails, etc.)

Important:
- You MUST execute the SKILL code after generating it
- You MUST run DRC to verify the layout is correct
- You MUST run PEX to extract and verify the capacitance
- Do not skip execution or verification steps
- If DRC fails, fix the issues and regenerate
- If PEX results show errors, iterate to improve the design"""
    },
    "capacitance_shape": {
        "description": "Unit and dummy capacitor design (capacitance + shape combinations)",
        "key_suffix": "cap_shape",
        "template": """Task: 
1. Phase 1: {task_desc} with target capacitance of {cap_value} fF
2. Phase 2: Generate dummy capacitor based on the unit capacitor

Design constraints:
- Target capacitance: {cap_value} fF
- Height constraint: less than 2 µm
- Low parasitic capacitance to ground required

Configuration:
- Technology: 28nm process node
- Library: LLM_Layout_Design
- Cell names: C_MAIN_{cap_str}fF_{shape_key} (unit), C_DUMMY_{cap_str}fF_{shape_key} (dummy)

Phase 1 - Unit Capacitor:
1. Generate unit {shape_desc} SKILL code with target capacitance of {cap_value} fF
2. Execute the generated SKILL code using run_il_file or run_il_with_screenshot tool
3. Run DRC verification using run_drc tool on the unit cell
4. Run PEX extraction using run_pex tool on the unit cell
5. Verify capacitance meets target ({cap_value} fF) and iterate if necessary

Phase 2 - Dummy Capacitor:
1. Generate dummy capacitor SKILL code based on the unit capacitor
2. Execute the generated SKILL code
3. Run DRC verification on the dummy cell
4. Run PEX extraction on the dummy cell

"""
    }
}

# Capacitance and shape configurations for capacitance_shape experiments
CAPACITANCE_VALUES = [0.1, 0.2, 0.3, 0.4, 0.5, 0.6, 0.7, 0.8, 0.9] + list(range(1, 11))
SHAPE_CONFIGS = [
    ("H", "H-shape capacitor", "Design unit H-shape capacitor"),
    ("H_shieldless", "H-shape capacitor without shield", "Design unit H-shape capacitor without shield"),
    ("I", "I-shape capacitor", "Design unit I-shape capacitor"),
    ("I_shield", "I-shape capacitor with shield", "Design unit I-shape capacitor with shield"),
    ("sandwich_simplified_h_notch", "Sandwich capacitor with simplified H-shape without notch", "Design unit Sandwich capacitor with simplified H-shape without notch"),
]

import subprocess
import sys
import os
import signal
import threading
from pathlib import Path
from datetime import datetime
import time

try:
    import openpyxl
except ImportError:
    print("Installing openpyxl...")
    os.system(f"{sys.executable} -m pip install openpyxl -q")
    import openpyxl

# ============================================================================
# Template Helper Functions
# ============================================================================

def get_available_templates():
    """Get list of available template types"""
    return list(PROMPT_TEMPLATES.keys())

def get_template_description(template_type):
    """Get description for a template type"""
    if template_type == "io_ring":
        return "IO Ring pad layout"
    return PROMPT_TEMPLATES.get(template_type, {}).get("description", "Unknown template")

# ============================================================================
# Core Functions
# ============================================================================

def load_sheet_names(excel_file):
    """Load all sheet names from Excel file"""
    try:
        wb = openpyxl.load_workbook(excel_file, read_only=True)
        sheet_names = wb.sheetnames
        wb.close()
        return sheet_names
    except Exception as e:
        print(f"Error loading Excel file {excel_file}: {e}")
        return []

def format_cap_str(cap):
    """Format capacitance for prompt key (0.1 -> 0_1, 1 -> 1)"""
    if cap < 1:
        return str(cap).replace('.', '_')
    else:
        return str(int(cap))

def generate_prompt_key(excel_name=None, sheet_name=None, prefix="", template_type="array", cap_value=None, shape_key=None):
    """Generate prompt key from parameters"""
    # Validate template type
    if template_type not in PROMPT_TEMPLATES:
        available = ", ".join(get_available_templates())
        raise ValueError(f"Unknown template type '{template_type}'. Available types: {available}")
    
    # Get key suffix from template registry
    key_suffix = PROMPT_TEMPLATES[template_type]["key_suffix"]
    
    if prefix:
        prefix = prefix.rstrip('_') + '_'  # Ensure single trailing underscore
    
    # Handle capacitance_shape template
    if template_type == "capacitance_shape":
        if cap_value is None or shape_key is None:
            raise ValueError("cap_value and shape_key are required for capacitance_shape template")
        cap_str = format_cap_str(cap_value)
        return f"{prefix}cap_{cap_str}fF_{shape_key}"
    
    # Handle CDAC templates (require excel_name and sheet_name)
    if excel_name is None or sheet_name is None:
        raise ValueError("excel_name and sheet_name are required for CDAC templates")
    
    # Convert to safe format for Virtuoso (replace spaces and '-' with '_')
    safe_excel_name = excel_name.replace(' ', '_').replace('-', '_')
    safe_sheet_name = sheet_name.replace(' ', '_').replace('-', '_')
    
    return f"{prefix}{key_suffix}_{safe_excel_name}_{safe_sheet_name}"

def generate_prompt_text(excel_file=None, sheet_name=None, prefix="", template_type="array", cap_value=None, shape_key=None, shape_desc=None, task_desc=None):
    """
    Generate prompt text dynamically based on template type and parameters
    
    Args:
        excel_file: Path to Excel file (required for CDAC templates)
        sheet_name: Name of the sheet (required for CDAC templates)
        prefix: Prefix for cell names (e.g., 'claude', 'gpt')
        template_type: Template type key (e.g., 'array', 'full', 'capacitance_shape')
        cap_value: Capacitance value in fF (required for capacitance_shape)
        shape_key: Shape key (required for capacitance_shape)
        shape_desc: Shape description (required for capacitance_shape)
        task_desc: Task description (required for capacitance_shape)
    
    Returns:
        Formatted prompt string
    
    Raises:
        ValueError: If template_type is not found or required parameters are missing
    """
    # Validate template type
    if template_type not in PROMPT_TEMPLATES:
        available = ", ".join(get_available_templates())
        raise ValueError(f"Unknown template type '{template_type}'. Available types: {available}")
    
    # Determine cell name prefix
    if prefix:
        cell_prefix = prefix.rstrip('_') + '_'
    else:
        cell_prefix = ""
    
    # Get template and key_suffix
    template = PROMPT_TEMPLATES[template_type]["template"]
    key_suffix = PROMPT_TEMPLATES[template_type]["key_suffix"]
    
    # Handle capacitance_shape template
    if template_type == "capacitance_shape":
        if cap_value is None or shape_key is None or shape_desc is None or task_desc is None:
            raise ValueError("cap_value, shape_key, shape_desc, and task_desc are required for capacitance_shape template")
        cap_str = format_cap_str(cap_value)
        prompt = template.format(
            cap_value=cap_value,
            cap_str=cap_str,
            shape_key=shape_key,
            shape_desc=shape_desc,
            task_desc=task_desc
        )
        return prompt
    
    # Handle CDAC templates
    if excel_file is None or sheet_name is None:
        raise ValueError("excel_file and sheet_name are required for CDAC templates")
    
    # Convert to safe format for Virtuoso (replace spaces and '-' with '_')
    safe_sheet_name = sheet_name.replace(' ', '_').replace('-', '_')
    
    # Format template with variables
    prompt = template.format(
        excel_file=excel_file,
        sheet_name=sheet_name,
        safe_sheet_name=safe_sheet_name,
        cell_prefix=cell_prefix,
        key_suffix=key_suffix
    )
    
    return prompt

def run_experiment(excel_file=None, sheet_name=None, prefix="", template_type="array", model_name=None, ramic_port=None, ramic_host=None, log_dir="logs/batch_experiments", batch_interrupted_flag=None, current_process_ref=None, cap_value=None, shape_key=None, shape_desc=None, task_desc=None, prompt_text=None, prompt_key=None):
    """Run a single experiment with dynamically generated prompt"""
    start_time = time.time()
    start_time_str = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    
    # Generate prompt key for logging/identification
    experiment_info = None
    if prompt_key is not None:
        # Use provided prompt key
        if template_type == "io_ring":
            experiment_info = f"Pad Layout: {sheet_name if sheet_name else 'N/A'}"
        elif template_type == "capacitance_shape":
            experiment_info = f"Capacitance: {cap_value} fF, Shape: {shape_key}"
        elif sheet_name:
            experiment_info = f"Sheet: {sheet_name}"
    elif template_type == "capacitance_shape":
        prompt_key = generate_prompt_key(prefix=prefix, template_type=template_type, cap_value=cap_value, shape_key=shape_key)
        experiment_info = f"Capacitance: {cap_value} fF, Shape: {shape_key}"
    elif template_type == "io_ring":
        # IO ring should provide prompt_key
        if prompt_key is None:
            raise ValueError("prompt_key is required for io_ring template type")
    else:
        excel_path = Path(excel_file)
        excel_name = excel_path.stem
        prompt_key = generate_prompt_key(excel_name, sheet_name, prefix, template_type)
        experiment_info = f"Sheet: {sheet_name}"
    
    print(f"\n{'='*80}")
    print(f"Running experiment: {prompt_key}")
    if experiment_info:
        print(f"{experiment_info}")
    print(f"Template type: {template_type} ({get_template_description(template_type)})")
    if ramic_port:
        print(f"Using RAMIC port: {ramic_port}")
    if ramic_host:
        print(f"Using RAMIC host: {ramic_host}")
    print(f"Start time: {start_time_str}")
    print(f"{'='*80}\n")
    
    # Create log directory if it doesn't exist
    os.makedirs(log_dir, exist_ok=True)
    
    # Generate prompt text dynamically (or use provided prompt_text)
    if prompt_text is None:
        if template_type == "capacitance_shape":
            prompt_text = generate_prompt_text(prefix=prefix, template_type=template_type, cap_value=cap_value, shape_key=shape_key, shape_desc=shape_desc, task_desc=task_desc)
        elif template_type == "io_ring":
            # IO ring experiments should provide prompt_text directly
            raise ValueError("prompt_text is required for io_ring template type")
        else:
            prompt_text = generate_prompt_text(excel_file, sheet_name, prefix, template_type)
    
    # Write prompt to temporary file (system will auto-load from user_prompt.txt)
    temp_prompt_file = Path("user_prompt.txt")
    try:
        with open(temp_prompt_file, 'w', encoding='utf-8') as f:
            f.write(prompt_text)
    except Exception as e:
        print(f"[ERROR] Failed to write prompt to temporary file: {e}")
        return {
            "prompt_key": prompt_key,
            "success": False,
            "elapsed_time": 0,
            "log_file": None,
            "error": f"Failed to write prompt file: {e}"
        }
    
    # Build command - use multi_agent_main.py as entry point (no --prompt, will auto-load from user_prompt.txt)
    cmd = [sys.executable, "-u", "src/multi_agent_main.py", "--log"]
    
    if model_name:
        cmd.extend(["--model-name", model_name])
    
    if ramic_host:
        cmd.extend(["--ramic-host", ramic_host])
    
    if ramic_port:
        cmd.extend(["--ramic-port", str(ramic_port)])
    
    # Log file for this experiment
    log_file = os.path.join(log_dir, f"{prompt_key}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.log")
    
    # Write experiment header to log file
    try:
        with open(log_file, 'w', encoding='utf-8') as f:
            f.write(f"{'='*80}\n")
            f.write(f"Experiment: {prompt_key}\n")
            f.write(f"Start time: {start_time_str}\n")
            if ramic_port:
                f.write(f"RAMIC port: {ramic_port}\n")
            if ramic_host:
                f.write(f"RAMIC host: {ramic_host}\n")
            f.write(f"{'='*80}\n\n")
    except Exception as e:
        print(f"[WARNING] Failed to write header to log file: {e}")
    success = False
    error_msg = None
    interrupted = False
    process = None  # Will be set later
    
    # Get current signal handlers (which should be the global ones) before defining new handler
    # We'll use a temporary handler to get the current one
    temp_handler = lambda s, f: None
    saved_sigint = signal.signal(signal.SIGINT, temp_handler)
    signal.signal(signal.SIGINT, saved_sigint)  # Restore immediately
    saved_sigterm = signal.signal(signal.SIGTERM, temp_handler)
    signal.signal(signal.SIGTERM, saved_sigterm)  # Restore immediately
    
    # Signal handler for graceful shutdown
    def signal_handler(signum, frame):
        nonlocal interrupted, error_msg, process
        # First, call the original handler (global handler) if it exists
        # This ensures the global handler sets the batch_interrupted flag
        if saved_sigint is not None and signum == signal.SIGINT:
            try:
                saved_sigint(signum, frame)
            except Exception:
                pass
        elif saved_sigterm is not None and signum == signal.SIGTERM:
            try:
                saved_sigterm(signum, frame)
            except Exception:
                pass
        
        # Then check if batch was interrupted (from global handler)
        if batch_interrupted_flag is not None and batch_interrupted_flag.get('flag', False):
            interrupted = True
            error_msg = "Batch interrupted"
            if process is not None:
                try:
                    if sys.platform.startswith('win'):
                        process.kill()
                    else:
                        os.killpg(os.getpgid(process.pid), signal.SIGKILL)
                except Exception:
                    pass
            return
        
        # If global handler hasn't been called yet, set the flag
        if batch_interrupted_flag is not None:
            batch_interrupted_flag['flag'] = True
        
        if not interrupted:  # Only print once
            interrupted = True
            error_msg = f"Experiment interrupted by signal {signum}"
            print(f"\n[WARNING] Received signal {signum}, terminating experiment and recording elapsed time...")
            # Try to terminate the process if it exists
            if process is not None:
                try:
                    if sys.platform.startswith('win'):
                        process.kill()
                    else:
                        os.killpg(os.getpgid(process.pid), signal.SIGKILL)
                except Exception:
                    pass
    
    # Register signal handlers (these will temporarily override global handlers)
    original_sigint = signal.signal(signal.SIGINT, signal_handler)
    original_sigterm = signal.signal(signal.SIGTERM, signal_handler)
    
    try:
        # Set environment to ensure unbuffered output
        env = os.environ.copy()
        env['PYTHONUNBUFFERED'] = '1'
        
        # Use Popen to start the process (append mode since we already wrote header)
        with open(log_file, 'a', encoding='utf-8') as f:
            if sys.platform.startswith('win'):
                process = subprocess.Popen(
                    cmd,
                    stdin=subprocess.PIPE,
                    stdout=subprocess.PIPE,
                    stderr=subprocess.STDOUT,
                    encoding='utf-8',
                    bufsize=0,  # Unbuffered
                    env=env,
                    cwd=Path(__file__).parent,
                    creationflags=subprocess.CREATE_NEW_PROCESS_GROUP
                )
            else:
                process = subprocess.Popen(
                    cmd,
                    stdin=subprocess.PIPE,
                    stdout=subprocess.PIPE,
                    stderr=subprocess.STDOUT,
                    encoding='utf-8',
                    bufsize=0,  # Unbuffered
                    env=env,
                    cwd=Path(__file__).parent,
                    preexec_fn=os.setsid
                )
            
            # Store process reference for global signal handler
            if current_process_ref is not None:
                current_process_ref['process'] = process
            
            # Indicators that the program is ready for input
            ready_indicators = [
                "[User prompt]:",
                "[User prompt]: ",
                "[User prompt]",
                "User prompt:",
            ]
            
            # Banner end indicators (alternative detection)
            banner_end_indicators = [
                "Reasoning Agent for Mixed-Signal IC",
                "╰────────────────────────────────",
            ]
            
            output_buffer = ""
            max_wait_time = 60  # Maximum wait time for initialization (seconds)
            init_start_time = time.time()
            banner_seen = False
            
            # Read output until we see the banner end or prompt
            while True:
                # Check if interrupted (local or batch-level)
                batch_int = batch_interrupted_flag is not None and batch_interrupted_flag.get('flag', False)
                if interrupted or batch_int:
                    if batch_int:
                        interrupted = True
                    print("[INFO] Experiment interrupted, stopping...")
                    break
                
                elapsed = time.time() - init_start_time
                if elapsed > max_wait_time:
                    print(f"[WARNING] Timeout waiting for initialization after {elapsed:.1f}s")
                    break
                
                if process.poll() is not None:
                    print(f"[WARNING] Process exited early with code {process.returncode}")
                    break
                
                try:
                    line = process.stdout.readline()
                    if line:
                        print(line, end='', flush=True)
                        f.write(line)
                        f.flush()
                        output_buffer += line
                        
                        # Check for banner end indicators
                        if any(indicator in line for indicator in banner_end_indicators):
                            banner_seen = True
                            break
                        
                        # Check for prompt indicators
                        if any(indicator in output_buffer for indicator in ready_indicators):
                            banner_seen = True
                            break
                except Exception as e:
                    time.sleep(0.1)
            
            # After banner, wait a bit more to ensure input() is ready
            if banner_seen and process.poll() is None and not interrupted:
                # Check batch interruption flag
                batch_int = batch_interrupted_flag is not None and batch_interrupted_flag.get('flag', False)
                if batch_int:
                    interrupted = True
                
                if not interrupted:
                    # Try to read one more line to see if there's a prompt
                    for _ in range(20):  # Wait up to 2 seconds
                        # Check interruption
                        batch_int = batch_interrupted_flag is not None and batch_interrupted_flag.get('flag', False)
                        if interrupted or batch_int or process.poll() is not None:
                            if batch_int:
                                interrupted = True
                            break
                        try:
                            if not sys.platform.startswith('win'):
                                import select
                                ready, _, _ = select.select([process.stdout], [], [], 0.1)
                                if ready:
                                    line = process.stdout.readline()
                                    if line:
                                        print(line, end='', flush=True)
                                        f.write(line)
                                        f.flush()
                                        output_buffer += line
                                        if any(ind in line for ind in ready_indicators):
                                            break
                            else:
                                # Windows: just wait a bit
                                time.sleep(0.5)
                                break
                        except:
                            time.sleep(0.1)
                    
                    # Wait a bit more to ensure input() is ready (only if not interrupted)
                    if not interrupted:
                        time.sleep(0.3)
            
            # Send "exit" command if process is still running (only if not interrupted)
            if not interrupted and process.poll() is None:
                try:
                    process.stdin.write("exit\n")
                    process.stdin.flush()
                except (BrokenPipeError, OSError) as e:
                    print(f"[WARNING] Failed to send exit command: {e}")
            
            # Continue reading output while process runs
            max_run_time = 3000  # 50 minutes max per experiment (expected ~40 minutes)
            run_start_time = time.time()
            
            while process.poll() is None and not interrupted and (batch_interrupted_flag is None or not batch_interrupted_flag.get('flag', False)):
                elapsed = time.time() - run_start_time
                if elapsed > max_run_time:
                    print(f"\n[WARNING] Timeout after {elapsed/60:.1f} minutes ({elapsed:.0f}s), force killing process...")
                    try:
                        if sys.platform.startswith('win'):
                            process.terminate()
                            time.sleep(2)
                            if process.poll() is None:
                                process.kill()
                        else:
                            os.killpg(os.getpgid(process.pid), signal.SIGTERM)
                            time.sleep(2)
                            if process.poll() is None:
                                os.killpg(os.getpgid(process.pid), signal.SIGKILL)
                    except Exception as e:
                        print(f"[WARNING] Error killing process: {e}")
                        try:
                            process.kill()
                        except:
                            pass
                    error_msg = f"Experiment timed out after {elapsed/60:.1f} minutes"
                    break
                
                # Read output
                try:
                    if not sys.platform.startswith('win'):
                        import select
                        # Use shorter timeout to check interruption more frequently
                        ready, _, _ = select.select([process.stdout], [], [], 0.2)
                        if ready:
                            line = process.stdout.readline()
                            if line:
                                print(line, end='', flush=True)
                                f.write(line)
                                f.flush()
                        # Check interruption even if no output ready
                        batch_int = batch_interrupted_flag is not None and batch_interrupted_flag.get('flag', False)
                        if batch_int:
                            interrupted = True
                            break
                    else:
                        # Windows: read line directly
                        line = process.stdout.readline()
                        if line:
                            print(line, end='', flush=True)
                            f.write(line)
                            f.flush()
                        else:
                            time.sleep(0.2)
                        # Check interruption
                        batch_int = batch_interrupted_flag is not None and batch_interrupted_flag.get('flag', False)
                        if batch_int:
                            interrupted = True
                            break
                except:
                    time.sleep(0.2)
                    # Check interruption even on exception
                    batch_int = batch_interrupted_flag is not None and batch_interrupted_flag.get('flag', False)
                    if batch_int:
                        interrupted = True
                        break
            
            # Read remaining output (only if not interrupted)
            if not interrupted:
                try:
                    for line in process.stdout:
                        # Check interruption during reading
                        batch_int = batch_interrupted_flag is not None and batch_interrupted_flag.get('flag', False)
                        if batch_int:
                            interrupted = True
                            break
                        print(line, end='', flush=True)
                        f.write(line)
                        f.flush()
                        if process.poll() is not None:
                            break
                except (BrokenPipeError, OSError):
                    pass
            
            # Check batch interruption flag
            if batch_interrupted_flag is not None and batch_interrupted_flag.get('flag', False):
                interrupted = True
            
            # Clear process reference
            if current_process_ref is not None:
                current_process_ref['process'] = None
            
            # Wait for process to finish (or kill if interrupted)
            if interrupted:
                # Force kill if interrupted
                try:
                    if sys.platform.startswith('win'):
                        process.kill()
                    else:
                        os.killpg(os.getpgid(process.pid), signal.SIGKILL)
                except Exception:
                    pass
                return_code = -1
            else:
                try:
                    return_code = process.wait(timeout=5)
                except subprocess.TimeoutExpired:
                    print("[WARNING] Process did not exit within timeout, force killing...")
                    try:
                        if sys.platform.startswith('win'):
                            process.kill()
                        else:
                            os.killpg(os.getpgid(process.pid), signal.SIGKILL)
                    except:
                        pass
                    return_code = -1
                    if not error_msg:
                        error_msg = "Process did not exit cleanly"
            
            success = (return_code == 0 or return_code is None)
            
            if return_code and return_code != 0 and not error_msg:
                error_msg = f"Experiment exited with return code {return_code}"
            
    except KeyboardInterrupt:
        interrupted = True
        error_msg = "Experiment interrupted by user (KeyboardInterrupt)"
        success = False
    except Exception as e:
        error_msg = f"Exception occurred: {str(e)}"
        success = False
    finally:
        # Clean up temporary prompt file
        try:
            temp_prompt_file = Path("user_prompt.txt")
            if temp_prompt_file.exists():
                temp_prompt_file.unlink()
        except Exception as e:
            print(f"[WARNING] Failed to clean up temporary prompt file: {e}")
        
        # Restore original signal handlers
        signal.signal(signal.SIGINT, original_sigint)
        signal.signal(signal.SIGTERM, original_sigterm)
        
        # Always record elapsed time, even if interrupted
        elapsed_time = time.time() - start_time
        
        # Update error message if interrupted
        if interrupted and not error_msg:
            error_msg = "Experiment interrupted"
        
        # Write summary to log file (always, even on interruption)
        try:
            with open(log_file, 'a', encoding='utf-8') as f:
                f.write(f"\n{'='*80}\n")
                if interrupted:
                    f.write(f"Experiment: {prompt_key} - ⚠ INTERRUPTED\n")
                else:
                    status = "✓ SUCCESS" if success else "✗ FAILED"
                    f.write(f"Experiment: {prompt_key} - {status}\n")
                f.write(f"Elapsed time: {elapsed_time/60:.2f} minutes ({elapsed_time:.1f} seconds)\n")
                f.write(f"Start time: {datetime.fromtimestamp(start_time).strftime('%Y-%m-%d %H:%M:%S')}\n")
                f.write(f"End time: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n")
                if error_msg:
                    f.write(f"Error: {error_msg}\n")
                f.write(f"{'='*80}\n")
        except Exception as e:
            print(f"[WARNING] Failed to write summary to log file: {e}")
        
        # Print summary
        print(f"\n{'='*80}")
        if interrupted:
            print(f"Experiment: {prompt_key} - ⚠ INTERRUPTED")
        else:
            status = "✓ SUCCESS" if success else "✗ FAILED"
            print(f"Experiment: {prompt_key} - {status}")
        print(f"Elapsed time: {elapsed_time/60:.2f} minutes ({elapsed_time:.1f} seconds)")
        if error_msg:
            print(f"Error: {error_msg}")
        print(f"Log file: {log_file}")
        print(f"{'='*80}\n")
    
    return {
        "prompt_key": prompt_key,
        "success": success,
        "elapsed_time": elapsed_time,
        "log_file": log_file,
        "error": error_msg
    }

def main():
    """Main function to run all batch experiments"""
    import argparse
    
    parser = argparse.ArgumentParser(
        description="Run batch experiments for CDAC or capacitance/shape combinations",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Examples:
  # CDAC experiments (default)
  python run_batch_experiments.py --prefix claude --model-name claude --template-type array
  python run_batch_experiments.py --prefix claude --model-name claude --template-type full
  
  # Capacitance/Shape experiments
  python run_batch_experiments.py --experiment-type capacitance_shape --prefix claude --model-name claude
        """
    )
    parser.add_argument(
        "--experiment-type",
        type=str,
        choices=["cdac", "capacitance_shape"],
        default="cdac",
        help="Experiment type: 'cdac' for CDAC array experiments, 'capacitance_shape' for capacitance/shape combinations. Default: 'cdac'"
    )
    parser.add_argument(
        "--excel-file",
        type=str,
        default="/home/lixintian/RAMIC_LXT/AMS-IO-Agent/excel/CDAC_3-8bit.xlsx",
        help="Path to Excel file (required for CDAC experiments, default: excel/CDAC_3-8bit.xlsx)"
    )
    parser.add_argument(
        "--prefix",
        type=str,
        default="",
        help="Prefix for prompt keys and cell names (e.g., 'claude', 'gpt', 'deepseek'). Default: no prefix"
    )
    parser.add_argument(
        "--template-type",
        type=str,
        choices=get_available_templates(),
        default="array",
        help="Prompt template type. Available: " + ", ".join([f"'{t}' ({get_template_description(t)})" for t in get_available_templates()]) + ". Default: 'array'"
    )
    parser.add_argument(
        "--model-name",
        type=str,
        default=None,
        help="Model name to use (e.g., deepseek, gpt-4o, claude). If not specified, uses default."
    )
    parser.add_argument(
        "--start-from",
        type=str,
        default=None,
        help="Start from a specific sheet name (useful for resuming)"
    )
    parser.add_argument(
        "--start-index",
        type=int,
        default=None,
        help="Start from a specific experiment index (1-based, e.g., 2 to start from the second experiment)"
    )
    parser.add_argument(
        "--stop-at",
        type=str,
        default=None,
        help="Stop at a specific sheet name"
    )
    parser.add_argument(
        "--stop-index",
        type=int,
        default=None,
        help="Stop at a specific experiment index (1-based, inclusive)"
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Only list the experiments that would be run, without executing them"
    )
    parser.add_argument(
        "--preview-prompt",
        type=str,
        nargs='?',
        const="first",
        help="Preview generated prompt(s). Use 'first' to preview first prompt, 'all' to preview all, or specify a sheet name. Example: --preview-prompt first"
    )
    parser.add_argument(
        "--ramic-host",
        type=str,
        default=None,
        help="RAMIC bridge host (default: from RB_HOST env var or 127.0.0.1)"
    )
    parser.add_argument(
        "--ramic-port-start",
        type=int,
        default=None,
        help="Starting RAMIC port number (each experiment will use port_start + index). If not specified, uses RB_PORT env var for all experiments."
    )
    parser.add_argument(
        "--ramic-port",
        type=int,
        default=None,
        help="RAMIC bridge port (used for all experiments if --ramic-port-start is not specified)"
    )
    
    args = parser.parse_args()
    
    # Validate experiment type and template type compatibility
    if args.experiment_type == "capacitance_shape":
        if args.template_type != "capacitance_shape":
            print(f"Warning: --experiment-type capacitance_shape requires --template-type capacitance_shape")
            print(f"Automatically setting --template-type to capacitance_shape")
            args.template_type = "capacitance_shape"
    elif args.experiment_type == "cdac":
        if args.template_type == "capacitance_shape":
            print(f"Error: --experiment-type cdac cannot use --template-type capacitance_shape")
            print(f"Available template types for CDAC: {', '.join([t for t in get_available_templates() if t != 'capacitance_shape'])}")
            return
    
    # Determine log directory based on prefix, template type, and experiment type
    if args.prefix:
        log_dir = f"logs/batch_{args.experiment_type}_{args.prefix}_{args.template_type}"
    else:
        log_dir = f"logs/batch_{args.experiment_type}_{args.template_type}"
    
    # Generate experiment list based on experiment type
    if args.experiment_type == "capacitance_shape":
        # Generate all capacitance and shape combinations
        experiments = []
        for cap in CAPACITANCE_VALUES:
            for shape_key, shape_desc, task_desc in SHAPE_CONFIGS:
                prompt_key = generate_prompt_key(prefix=args.prefix, template_type=args.template_type, cap_value=cap, shape_key=shape_key)
                experiments.append({
                    "prompt_key": prompt_key,
                    "cap_value": cap,
                    "shape_key": shape_key,
                    "shape_desc": shape_desc,
                    "task_desc": task_desc
                })
        
        # Sort by capacitance then shape
        experiments.sort(key=lambda x: (x["cap_value"], x["shape_key"]))
        
        print(f"Generating {len(experiments)} capacitance/shape experiments:")
        print(f"  Capacitance values: {len(CAPACITANCE_VALUES)} ({min(CAPACITANCE_VALUES)} - {max(CAPACITANCE_VALUES)} fF)")
        print(f"  Shapes: {len(SHAPE_CONFIGS)} ({', '.join([s[0] for s in SHAPE_CONFIGS])})")
        if args.prefix:
            print(f"Using prefix: '{args.prefix}'")
        print(f"Template type: {args.template_type} ({get_template_description(args.template_type)})")
        if args.model_name:
            print(f"Using model: '{args.model_name}'")
        print(f"Log directory: {log_dir}")
        for i, exp in enumerate(experiments, 1):
            print(f"  {i}. {exp['cap_value']} fF, {exp['shape_key']} -> {exp['prompt_key']}")
        
        # Filter by start/stop index
        if args.start_index is not None:
            if args.start_index < 1:
                print(f"Warning: Start index must be >= 1, ignoring --start-index")
            elif args.start_index > len(experiments):
                print(f"Warning: Start index {args.start_index} exceeds total experiments ({len(experiments)}), starting from beginning")
            else:
                experiments = experiments[args.start_index - 1:]
                print(f"\nStarting from experiment index: {args.start_index}")
        
        if args.stop_index is not None:
            if args.stop_index < 1:
                print(f"Warning: Stop index must be >= 1, ignoring --stop-index")
            elif args.stop_index > len(experiments):
                print(f"Warning: Stop index {args.stop_index} exceeds remaining experiments ({len(experiments)}), running all remaining")
            else:
                experiments = experiments[:args.stop_index]
                print(f"Stopping at experiment index: {args.stop_index}")
        
        if args.dry_run:
            print("\n[DRY RUN] Would run the following experiments:")
            for exp in experiments:
                print(f"  - {exp['cap_value']} fF, {exp['shape_key']} -> {exp['prompt_key']}")
            return
        
        if args.preview_prompt:
            print("\n" + "="*80)
            print("PROMPT PREVIEW")
            print("="*80)
            
            if args.preview_prompt.lower() == "first":
                if experiments:
                    exp = experiments[0]
                    prompt_text = generate_prompt_text(prefix=args.prefix, template_type=args.template_type, 
                                                       cap_value=exp["cap_value"], shape_key=exp["shape_key"],
                                                       shape_desc=exp["shape_desc"], task_desc=exp["task_desc"])
                    print(f"\nCapacitance: {exp['cap_value']} fF, Shape: {exp['shape_key']}")
                    print(f"Prompt Key: {exp['prompt_key']}")
                    print("-"*80)
                    print(prompt_text)
                    print("="*80)
            elif args.preview_prompt.lower() == "all":
                for i, exp in enumerate(experiments, 1):
                    prompt_text = generate_prompt_text(prefix=args.prefix, template_type=args.template_type,
                                                       cap_value=exp["cap_value"], shape_key=exp["shape_key"],
                                                       shape_desc=exp["shape_desc"], task_desc=exp["task_desc"])
                    print(f"\n[{i}/{len(experiments)}] {exp['cap_value']} fF, {exp['shape_key']}")
                    print(f"Prompt Key: {exp['prompt_key']}")
                    print("-"*80)
                    print(prompt_text)
                    if i < len(experiments):
                        print("\n" + "="*80)
                print("="*80)
            return
        
        # Store experiments for later execution
        experiment_list = experiments
        
    else:  # CDAC experiments
        # Get Excel file name without extension for prompt key generation
        excel_path = Path(args.excel_file)
        excel_name = excel_path.stem  # e.g., "CDAC_3-8bit"
        
        # Load sheet names from Excel
        print(f"Loading sheet names from {args.excel_file}...")
        sheet_names = load_sheet_names(args.excel_file)
        
        if not sheet_names:
            print(f"No sheets found in {args.excel_file}")
            return
        
        # Generate prompt keys from sheet names (preserve Excel sheet order)
        sheet_to_key = {}
        for sheet_name in sheet_names:
            prompt_key = generate_prompt_key(excel_name, sheet_name, args.prefix, args.template_type)
            sheet_to_key[sheet_name] = prompt_key
        
        # Create list of (sheet_name, prompt_key) tuples in Excel order (no sorting)
        sheet_key_pairs = [(s, sheet_to_key[s]) for s in sheet_names]
        
        # Extract lists (preserving Excel order)
        sorted_sheet_names = [pair[0] for pair in sheet_key_pairs]
        prompt_keys = [pair[1] for pair in sheet_key_pairs]
        
        print(f"Found {len(sheet_names)} sheets, generating {len(prompt_keys)} experiments:")
        if args.prefix:
            print(f"Using prefix: '{args.prefix}'")
        print(f"Template type: {args.template_type} ({get_template_description(args.template_type)})")
        if args.model_name:
            print(f"Using model: '{args.model_name}'")
        print(f"Log directory: {log_dir}")
        for i, (sheet_name, prompt_key) in enumerate(sheet_key_pairs, 1):
            print(f"  {i}. Sheet: {sheet_name} -> Prompt key: {prompt_key}")
        
        # Filter by start/stop if specified (by sheet name or index)
        if args.start_index is not None:
            # Start from index (1-based, convert to 0-based)
            if args.start_index < 1:
                print(f"Warning: Start index must be >= 1, ignoring --start-index")
            elif args.start_index > len(sheet_key_pairs):
                print(f"Warning: Start index {args.start_index} exceeds total experiments ({len(sheet_key_pairs)}), starting from beginning")
            else:
                start_idx = args.start_index - 1  # Convert to 0-based
                sheet_key_pairs = sheet_key_pairs[start_idx:]
                sorted_sheet_names = [pair[0] for pair in sheet_key_pairs]
                prompt_keys = [pair[1] for pair in sheet_key_pairs]
                print(f"\nStarting from experiment index: {args.start_index} ({sheet_key_pairs[0][0]})")
        elif args.start_from:
            if args.start_from in sorted_sheet_names:
                start_idx = sorted_sheet_names.index(args.start_from)
                sorted_sheet_names = sorted_sheet_names[start_idx:]
                prompt_keys = [sheet_to_key[s] for s in sorted_sheet_names]
                sheet_key_pairs = [(s, sheet_to_key[s]) for s in sorted_sheet_names]
                print(f"\nStarting from sheet: {args.start_from}")
            else:
                print(f"Warning: Sheet '{args.start_from}' not found in Excel file")
        
        if args.stop_index is not None:
            # Stop at index (1-based, convert to 0-based)
            if args.stop_index < 1:
                print(f"Warning: Stop index must be >= 1, ignoring --stop-index")
            elif args.stop_index > len(sheet_key_pairs):
                print(f"Warning: Stop index {args.stop_index} exceeds remaining experiments ({len(sheet_key_pairs)}), running all remaining")
            else:
                stop_idx = args.stop_index  # Keep 1-based for slicing (exclusive)
                sheet_key_pairs = sheet_key_pairs[:stop_idx]
                sorted_sheet_names = [pair[0] for pair in sheet_key_pairs]
                prompt_keys = [pair[1] for pair in sheet_key_pairs]
                print(f"Stopping at experiment index: {args.stop_index} ({sheet_key_pairs[-1][0]})")
        elif args.stop_at:
            if args.stop_at in sorted_sheet_names:
                stop_idx = sorted_sheet_names.index(args.stop_at)
                sorted_sheet_names = sorted_sheet_names[:stop_idx + 1]
                prompt_keys = [sheet_to_key[s] for s in sorted_sheet_names]
                sheet_key_pairs = [(s, sheet_to_key[s]) for s in sorted_sheet_names]
                print(f"Stopping at sheet: {args.stop_at}")
            else:
                print(f"Warning: Sheet '{args.stop_at}' not found in Excel file")
        
        if args.dry_run:
            print("\n[DRY RUN] Would run the following experiments:")
            for sheet_name, prompt_key in sheet_key_pairs:
                print(f"  - Sheet: {sheet_name} -> {prompt_key}")
            return
        
        if args.preview_prompt:
            print("\n" + "="*80)
            print("PROMPT PREVIEW")
            print("="*80)
            
            if args.preview_prompt.lower() == "first":
                # Preview first prompt
                if sheet_key_pairs:
                    sheet_name, prompt_key = sheet_key_pairs[0]
                    prompt_text = generate_prompt_text(args.excel_file, sheet_name, args.prefix, args.template_type)
                    print(f"\nSheet: {sheet_name}")
                    print(f"Prompt Key: {prompt_key}")
                    print("-"*80)
                    print(prompt_text)
                    print("="*80)
                else:
                    print("No experiments to preview.")
            elif args.preview_prompt.lower() == "all":
                # Preview all prompts
                for i, (sheet_name, prompt_key) in enumerate(sheet_key_pairs, 1):
                    prompt_text = generate_prompt_text(args.excel_file, sheet_name, args.prefix, args.template_type)
                    print(f"\n[{i}/{len(sheet_key_pairs)}] Sheet: {sheet_name}")
                    print(f"Prompt Key: {prompt_key}")
                    print("-"*80)
                    print(prompt_text)
                    if i < len(sheet_key_pairs):
                        print("\n" + "="*80)
                print("="*80)
            else:
                # Preview specific sheet
                sheet_name = args.preview_prompt
                found = False
                for s_name, prompt_key in sheet_key_pairs:
                    if s_name == sheet_name:
                        prompt_text = generate_prompt_text(args.excel_file, sheet_name, args.prefix, args.template_type)
                        print(f"\nSheet: {sheet_name}")
                        print(f"Prompt Key: {prompt_key}")
                        print("-"*80)
                        print(prompt_text)
                        print("="*80)
                        found = True
                        break
                if not found:
                    print(f"Error: Sheet '{sheet_name}' not found in Excel file.")
                    print(f"Available sheets: {', '.join([s for s, _ in sheet_key_pairs])}")
            return
        
        # Store experiments for later execution (CDAC format)
        experiment_list = [{"sheet_name": s, "prompt_key": k} for s, k in sheet_key_pairs]
    
    # Show summary before running
    print(f"\n{'='*80}")
    print(f"About to run {len(experiment_list)} experiments")
    if args.model_name:
        print(f"Model: {args.model_name}")
    print(f"Each experiment will timeout after 50 minutes if not completed")
    print(f"All experiments will run automatically without user intervention")
    print(f"{'='*80}")
    print("Starting batch execution in 3 seconds...")
    time.sleep(3)
    
    # Global flag for batch interruption (use dict for shared state)
    batch_interrupted = {'flag': False}
    current_process = {'process': None}  # Store current process for immediate termination
    
    # Global signal handler for batch interruption
    def batch_signal_handler(signum, frame):
        if not batch_interrupted['flag']:
            batch_interrupted['flag'] = True
            print(f"\n\n{'='*80}")
            print(f"[BATCH INTERRUPTED] Received signal {signum}")
            print(f"Terminating current experiment and stopping batch...")
            print(f"{'='*80}\n")
            # Immediately kill current process if it exists
            if current_process['process'] is not None:
                try:
                    if sys.platform.startswith('win'):
                        current_process['process'].kill()
                    else:
                        os.killpg(os.getpgid(current_process['process'].pid), signal.SIGKILL)
                except Exception:
                    pass
    
    # Register global signal handlers (will override experiment-level handlers)
    original_sigint = signal.signal(signal.SIGINT, batch_signal_handler)
    original_sigterm = signal.signal(signal.SIGTERM, batch_signal_handler)
    
    # Run experiments
    results = []
    total_start_time = time.time()
    
    print(f"Press Ctrl+C once to stop the entire batch immediately")
    
    try:
        for i, exp in enumerate(experiment_list, 1):
            # Check if batch was interrupted
            if batch_interrupted['flag']:
                print(f"\n[BATCH STOPPED] Stopping batch execution after {len(results)} experiments")
                break
            
            prompt_key = exp["prompt_key"]
            
            # Determine RAMIC port for this experiment
            ramic_port = None
            if args.ramic_port_start is not None:
                # Use sequential ports starting from ramic_port_start
                ramic_port = args.ramic_port_start + (i - 1)
            elif args.ramic_port is not None:
                # Use the same port for all experiments
                ramic_port = args.ramic_port
            
            # Run experiment based on type
            if args.experiment_type == "capacitance_shape":
                print(f"\n[{i}/{len(experiment_list)}] Processing: {prompt_key} ({exp['cap_value']} fF, {exp['shape_key']})")
                result = run_experiment(
                    prefix=args.prefix,
                    template_type=args.template_type,
                    model_name=args.model_name, 
                    ramic_port=ramic_port, 
                    ramic_host=args.ramic_host,
                    log_dir=log_dir,
                    batch_interrupted_flag=batch_interrupted,
                    current_process_ref=current_process,
                    cap_value=exp["cap_value"],
                    shape_key=exp["shape_key"],
                    shape_desc=exp["shape_desc"],
                    task_desc=exp["task_desc"]
                )
            else:  # CDAC experiments
                print(f"\n[{i}/{len(experiment_list)}] Processing: {prompt_key} (Sheet: {exp['sheet_name']})")
                result = run_experiment(
                    args.excel_file,
                    exp["sheet_name"],
                    prefix=args.prefix,
                    template_type=args.template_type,
                    model_name=args.model_name, 
                    ramic_port=ramic_port, 
                    ramic_host=args.ramic_host,
                    log_dir=log_dir,
                    batch_interrupted_flag=batch_interrupted,
                    current_process_ref=current_process
                )
            results.append(result)
            
            # Check again after experiment (in case it was interrupted during execution)
            if batch_interrupted['flag']:
                print(f"\n[BATCH STOPPED] Stopping batch execution after {len(results)} experiments")
                break
            
            # If failed, automatically continue to next experiment (no user interaction)
            if not result["success"]:
                is_timeout = result.get("error", "").startswith("Experiment timed out")
                if is_timeout:
                    print(f"\n⚠️  Experiment '{prompt_key}' timed out after 50 minutes, automatically continuing to next experiment...")
                else:
                    print(f"\n⚠️  Experiment '{prompt_key}' failed: {result.get('error', 'Unknown error')}")
                    print(f"Automatically continuing to next experiment...")
    except KeyboardInterrupt:
        batch_interrupted['flag'] = True
        print(f"\n[BATCH INTERRUPTED] Stopping batch execution...")
    finally:
        # Restore original signal handlers
        signal.signal(signal.SIGINT, original_sigint)
        signal.signal(signal.SIGTERM, original_sigterm)
    
    # Print summary
    total_time = time.time() - total_start_time
    successful = sum(1 for r in results if r["success"])
    failed = len(results) - successful
    
    print(f"\n{'='*80}")
    if batch_interrupted['flag']:
        print("BATCH EXPERIMENT SUMMARY (INTERRUPTED)")
    else:
        print("BATCH EXPERIMENT SUMMARY")
    print(f"{'='*80}")
    print(f"Total experiments completed: {len(results)}/{len(experiment_list)}")
    if batch_interrupted['flag']:
        print(f"Remaining experiments: {len(experiment_list) - len(results)}")
    print(f"Successful: {successful}")
    print(f"Failed: {failed}")
    print(f"Total time: {total_time/3600:.2f} hours ({total_time/60:.2f} minutes)")
    print(f"\nResults:")
    for result in results:
        status = "✓" if result["success"] else "✗"
        print(f"  {status} {result['prompt_key']} ({result['elapsed_time']/60:.2f} min)")
        if not result["success"] and result["error"]:
            print(f"      Error: {result['error']}")
    
    if batch_interrupted['flag'] and len(results) < len(experiment_list):
        print(f"\nInterrupted experiments (not run):")
        for exp in experiment_list[len(results):]:
            print(f"  - {exp['prompt_key']}")
    
    print(f"{'='*80}\n")
    
    # Save summary to file
    summary_file = os.path.join(log_dir, f"summary_{datetime.now().strftime('%Y%m%d_%H%M%S')}.txt")
    os.makedirs(log_dir, exist_ok=True)
    with open(summary_file, 'w', encoding='utf-8') as f:
        f.write("BATCH EXPERIMENT SUMMARY\n")
        f.write("="*80 + "\n")
        f.write(f"Experiment type: {args.experiment_type}\n")
        f.write(f"Template type: {args.template_type}\n")
        f.write(f"Total experiments: {len(results)}/{len(experiment_list)}\n")
        f.write(f"Successful: {successful}\n")
        f.write(f"Failed: {failed}\n")
        f.write(f"Total time: {total_time/3600:.2f} hours ({total_time/60:.2f} minutes)\n\n")
        f.write("Results:\n")
        for result in results:
            status = "✓" if result["success"] else "✗"
            f.write(f"  {status} {result['prompt_key']} ({result['elapsed_time']/60:.2f} min)\n")
            if not result["success"] and result["error"]:
                f.write(f"      Error: {result['error']}\n")
            f.write(f"      Log: {result['log_file']}\n")
    
    print(f"Summary saved to: {summary_file}")

if __name__ == "__main__":
    main()

